"""Verify the 3 near-real-world Vietnamese .xlsx benchmark workbooks in public/samples/."""

from __future__ import annotations

import json
from pathlib import Path

from openpyxl import load_workbook

from greenlogix_api.generate_realworld_samples import generate_all


def test_realworld_xlsx_samples_exist_and_match_schema() -> None:
    repo_root = Path(__file__).resolve().parents[3]
    samples_dir = repo_root / "apps" / "landing" / "public" / "samples"
    manifest_path = samples_dir / "samples_manifest.json"
    if not manifest_path.exists():
        generate_all(samples_dir)

    manifest = json.loads(manifest_path.read_text(encoding="utf-8"))
    expected_counts = {
        "EcoMiles_Mau_FMCG_80_Don_TPHCM.xlsx": 80,
        "EcoMiles_Mau_Chuoi_Lanh_120_Don_Cam_Tai_QD23.xlsx": 120,
        "EcoMiles_Mau_ERP_500_Don_7_Ngay.xlsx": 500,
    }
    assert set(manifest.keys()) == set(expected_counts.keys())

    for fn, count in expected_counts.items():
        xlsx_path = samples_dir / fn
        assert xlsx_path.exists(), f"Missing sample workbook: {fn}"
        wb = load_workbook(xlsx_path, read_only=True, data_only=True)
        assert "orders" in wb.sheetnames
        ws = wb["orders"]
        rows = list(ws.iter_rows(values_only=True))
        assert len(rows) == count + 1
        header = [str(c) for c in rows[0]]
        assert "external_id" in header
        assert "address" in header
        assert "kg" in header
        wb.close()
