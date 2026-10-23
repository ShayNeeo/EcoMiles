"""C-03: ingest at scale, idempotent re-upload, rejected rows, generator determinism."""

from __future__ import annotations

import hashlib
import json
import logging
import re
import sqlite3
from datetime import datetime
from io import BytesIO
from pathlib import Path
from zipfile import ZipFile

from openpyxl import Workbook, load_workbook
from sqlalchemy.exc import OperationalError
from sqlmodel import Session

from greenlogix_api import db as dbmod
from greenlogix_api.generate_synthetic import (
    COLUMNS,
    FIXED_STAMP,
    generate,
    order_rows,
    xlsx_bytes,
)
from greenlogix_api.ingest_xlsx import (
    DEFAULT_MAX_DATA_ROWS,
    DEFAULT_MAX_UPLOAD_BYTES,
    MAX_UPLOAD_BYTES_CEILING,
    IngestStream,
    resolve_limits,
)
from greenlogix_api.schemas import CARGO_TYPES, LOCAL_TIME_PATTERN
from greenlogix_api.seed import DISTRICTS

AUTH = {"Authorization": "Bearer DEMO"}
XLSX_CT = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
LEGACY_HEADERS = COLUMNS[1:]


def _post(client, payload: bytes, name: str = "orders.xlsx"):
    return client.post(
        "/orders/import", headers=AUTH, files={"file": (name, payload, XLSX_CT)}
    )


def _xlsx(headers: list[str], rows: list[list[object]]) -> bytes:
    wb = Workbook()
    ws = wb.active
    assert ws is not None
    ws.append(headers)
    for row in rows:
        ws.append(row)
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _generated(rows: int, seed: int, out_dir: Path) -> bytes:
    meta = generate(rows, seed=seed, out_dir=out_dir)
    return (out_dir / str(meta["xlsx"])).read_bytes()


def _legacy_xlsx(rows: list[list[object]]) -> bytes:
    return _xlsx(LEGACY_HEADERS, rows)


def test_demo_seed_file_imports_unchanged(demo_client) -> None:
    """The shipped 80-order demo file (no external column) still imports as before."""
    # the canonical demo workbook, not the per-test SEED_DIR the fixture patches
    payload = (dbmod.API_ROOT / "data" / "seed" / "hcmc_80_orders.xlsx").read_bytes()
    first = _post(demo_client, payload, "hcmc_80_orders.xlsx").json()
    assert first["imported"] == 80
    assert first["rows_seen"] == 80
    assert first["rows_imported"] == 80
    assert first["skipped_duplicates"] == 0
    assert first["dedupe"] is False
    assert first["errors"] == []
    # No external key => a re-upload duplicates, exactly as before C-03.
    assert _post(demo_client, payload, "hcmc_80_orders.xlsx").json()["rows_imported"] == 80
    assert len(demo_client.get("/orders", headers=AUTH).json()) == 160


def test_imports_over_500_rows(demo_client, tmp_path) -> None:
    payload = _generated(600, 7, tmp_path)
    res = _post(demo_client, payload, "sim-600.xlsx")
    assert res.status_code == 200
    body = res.json()
    assert body["rows_imported"] == 600
    assert body["imported"] == 600
    assert body["rows_seen"] == 600
    assert body["rows_rejected"] == 0
    assert body["skipped_duplicates"] == 0
    assert body["dedupe"] is True
    assert body["limits"] == {
        "max_upload_bytes": DEFAULT_MAX_UPLOAD_BYTES,
        "max_data_rows": DEFAULT_MAX_DATA_ROWS,
    }
    assert len(demo_client.get("/orders", headers=AUTH).json()) == 600


def test_reupload_skips_duplicates(demo_client, tmp_path) -> None:
    payload = _generated(120, 11, tmp_path)
    first = _post(demo_client, payload).json()
    assert first["rows_imported"] == 120
    assert first["skipped_duplicates"] == 0
    assert first["errors"] == []

    second = _post(demo_client, payload).json()
    assert second["rows_imported"] == 0
    assert second["imported"] == 0
    assert second["rows_seen"] == 120
    assert second["skipped_duplicates"] == 120
    assert second["rows_rejected"] == 0
    assert second["errors"] == []
    assert second["dedupe"] is True
    assert len(demo_client.get("/orders", headers=AUTH).json()) == 120

    # A duplicate inside one file is skipped too, without an error entry.
    rows = order_rows(2, 3)
    rows[1][0] = rows[0][0]
    within = _post(demo_client, xlsx_bytes(rows), "dupes.xlsx").json()
    assert within["rows_seen"] == 2
    assert within["rows_imported"] == 1
    assert within["skipped_duplicates"] == 1
    assert within["errors"] == []


def test_bad_row_reported_and_good_rows_imported(demo_client, tmp_path) -> None:
    rows = order_rows(3, 5)
    rows[1][2] = None  # lat
    res = _post(demo_client, xlsx_bytes(rows), "bad-row.xlsx")
    assert res.status_code == 200
    body = res.json()
    assert body["rows_seen"] == 3
    assert body["rows_imported"] == 2
    assert body["rows_rejected"] == 1
    assert body["skipped_duplicates"] == 0
    assert body["errors"] == [{"excel_row": 3, "field": "lat", "message": "missing lat"}]
    assert len(demo_client.get("/orders", headers=AUTH).json()) == 2


def test_file_without_external_id_keeps_old_behaviour(demo_client) -> None:
    row = ["Q1", 10.776, 106.7, "KH", "0900000001", 12, "08:00", "10:00", "thuong", "ghi"]
    payload = _legacy_xlsx([row, [*row[:5], 7, "09:00", "11:00", "lanh", ""]])
    first = _post(demo_client, payload).json()
    assert first["dedupe"] is False
    assert first["rows_imported"] == 2
    assert first["imported"] == 2
    # No external key => nothing to dedupe on, exactly as before C-03.
    second = _post(demo_client, payload).json()
    assert second["dedupe"] is False
    assert second["rows_imported"] == 2
    assert second["skipped_duplicates"] == 0
    assert len(demo_client.get("/orders", headers=AUTH).json()) == 4


def test_limits_are_env_overridable(demo_client, monkeypatch, tmp_path) -> None:
    payload = _generated(6, 3, tmp_path)
    monkeypatch.setenv("ECOMILES_MAX_DATA_ROWS", "5")
    assert resolve_limits() == (DEFAULT_MAX_UPLOAD_BYTES, 5)
    assert _post(demo_client, payload).status_code == 400
    monkeypatch.setenv("ECOMILES_MAX_UPLOAD_BYTES", "1000")
    assert resolve_limits() == (1000, 5)
    assert _post(demo_client, payload).status_code == 400
    monkeypatch.setenv("ECOMILES_MAX_DATA_ROWS", "not-an-int")
    assert resolve_limits() == (1000, DEFAULT_MAX_DATA_ROWS)
    assert demo_client.get("/orders", headers=AUTH).json() == []


def test_row_cap_rolls_back_flushed_batches(demo_client, monkeypatch, tmp_path) -> None:
    """The 500-row batches share one transaction: a late cap breach inserts nothing."""
    monkeypatch.setenv("ECOMILES_MAX_DATA_ROWS", "600")
    res = _post(demo_client, _generated(700, 5, tmp_path), "too-many.xlsx")
    assert res.status_code == 400
    assert res.json()["detail"] == "too many rows"
    assert demo_client.get("/orders", headers=AUTH).json() == []


def test_generator_output_is_deterministic(tmp_path) -> None:
    first = generate(80, seed=99, out_dir=tmp_path / "a")
    second = generate(80, seed=99, out_dir=tmp_path / "b")
    other = generate(80, seed=100, out_dir=tmp_path / "c")
    digests = {d: hashlib.sha256((tmp_path / d / first["xlsx"]).read_bytes()).hexdigest() for d in "abc"}
    assert digests["a"] == digests["b"] == first["sha256"] == second["sha256"]
    assert digests["c"] == other["sha256"]
    assert digests["a"] != digests["c"]
    # Reproducible across processes only because both workbook stamps are pinned
    # (openpyxl re-stamps ``modified`` with the wall clock at save time).
    with ZipFile(BytesIO((tmp_path / "a" / str(first["xlsx"])).read_bytes())) as archive:
        core = archive.read("docProps/core.xml").decode("utf-8")
    assert core.count(FIXED_STAMP) == 2


def test_generator_metadata_and_lat_lng(tmp_path) -> None:
    meta = generate(80, seed=20261018, out_dir=tmp_path)
    assert meta["label"] == "sim-hcmc-80"
    assert meta["kind"] == "simulated"
    assert meta["seed"] == 20261018
    assert meta["rows"] == 80
    assert meta["columns"] == COLUMNS
    assert meta["xlsx"] == "synthetic_80_orders.xlsx"
    datetime.fromisoformat(str(meta["generated_at"]))
    saved = json.loads((tmp_path / "dataset_meta.json").read_text(encoding="utf-8"))
    assert saved == meta

    wb = load_workbook(tmp_path / "synthetic_80_orders.xlsx", read_only=True, data_only=True)
    ws = wb.active
    assert ws is not None
    rows = list(ws.iter_rows(min_row=2, values_only=True))
    wb.close()
    assert len(rows) == 80
    lat_col, lng_col = COLUMNS.index("lat"), COLUMNS.index("lng")
    for row in rows:
        lat, lng = row[lat_col], row[lng_col]
        assert isinstance(lat, float) and isinstance(lng, float)
        assert any(
            abs(lat - district_lat) < 0.01 and abs(lng - district_lng) < 0.01
            for _, district_lat, district_lng in DISTRICTS
        )


def test_unrepresentable_fields_are_rejected_or_normalised(demo_client, tmp_path) -> None:
    """F1: nothing the response model would reject may reach the database."""
    rows = order_rows(4, 21)
    rows[1][7] = None  # window_start blank (ordinary spreadsheet output)
    rows[2][7] = "8:00"  # unpadded hour: normalised, not dropped
    rows[3][9] = "hang_de"  # cargo_type outside the enum
    res = _post(demo_client, xlsx_bytes(rows), "bad-fields.xlsx")
    assert res.status_code == 200
    body = res.json()
    assert body["rows_seen"] == 4
    assert body["rows_imported"] == 2
    assert body["rows_rejected"] == 2
    assert body["errors"] == [
        {"excel_row": 3, "field": "window_start", "message": "missing window_start"},
        {"excel_row": 5, "field": "cargo_type", "message": "unknown cargo_type"},
    ]
    listed = demo_client.get("/orders", headers=AUTH)
    assert listed.status_code == 200
    orders = listed.json()
    assert len(orders) == 2
    for order in orders:
        assert re.fullmatch(LOCAL_TIME_PATTERN, order["window_start"])
        assert re.fullmatch(LOCAL_TIME_PATTERN, order["window_end"])
        assert order["cargo_type"] in CARGO_TYPES
    assert any(o["window_start"] == "08:00" for o in orders)


def test_file_without_window_columns_imports_nothing(demo_client) -> None:
    """A workbook with no window column cannot be stored: OrderOut needs HH:MM."""
    payload = _xlsx(
        ["address", "lat", "lng", "receiver", "phone", "kg", "cargo_type", "notes"],
        [["Q1", 10.776, 106.7, "KH", "0900000009", 5, "thuong", ""]],
    )
    body = _post(demo_client, payload, "no-windows.xlsx").json()
    assert body["rows_imported"] == 0
    assert body["errors"][0]["field"] == "window_start"
    assert demo_client.get("/orders", headers=AUTH).status_code == 200


def test_oversize_upload_rejected_and_ceiling_clamped(demo_client, monkeypatch, tmp_path) -> None:
    """F2: the cap is checked before reading, and absurd env values are clamped."""
    payload = _generated(20, 13, tmp_path)
    assert len(payload) > 2_000
    monkeypatch.setenv("ECOMILES_MAX_UPLOAD_BYTES", "2000")
    res = _post(demo_client, payload, "too-big.xlsx")
    assert res.status_code == 400
    assert res.json()["detail"] == "file too large"
    assert demo_client.get("/orders", headers=AUTH).json() == []
    monkeypatch.setenv("ECOMILES_MAX_UPLOAD_BYTES", str(10**12))
    assert resolve_limits()[0] == MAX_UPLOAD_BYTES_CEILING
    monkeypatch.setenv("ECOMILES_MAX_UPLOAD_BYTES", str(MAX_UPLOAD_BYTES_CEILING + 1))
    assert resolve_limits()[0] == MAX_UPLOAD_BYTES_CEILING
    assert MAX_UPLOAD_BYTES_CEILING == 256 * 1024 * 1024


def test_phone_query_is_not_logged_verbatim(demo_client, caplog) -> None:
    """F3: the dispatcher query can be a phone number, so it never reaches logs."""
    caplog.set_level(logging.INFO, logger="greenlogix")
    phone = "0901234567"
    assert demo_client.get(f"/orders?q={phone}", headers=AUTH).status_code == 200
    assert "path=/orders" in caplog.text
    assert phone not in caplog.text
    assert "09***" in caplog.text


def test_database_failure_is_500_not_invalid_xlsx(monkeypatch, tmp_path) -> None:
    """M-7: a commit/flush failure must not be reported as a bad spreadsheet."""
    monkeypatch.setenv("GREENLOGIX_DEMO", "1")
    from fastapi.testclient import TestClient

    from greenlogix_api.main import app

    rolled_back: list[str] = []
    real_rollback = Session.rollback

    def _failing_commit(_self: Session) -> None:
        raise OperationalError("COMMIT", {}, Exception("disk I/O error"))

    def _spy_rollback(self: Session) -> None:
        rolled_back.append("rollback")
        real_rollback(self)

    monkeypatch.setattr(Session, "commit", _failing_commit)
    monkeypatch.setattr(Session, "rollback", _spy_rollback)
    payload = _generated(600, 17, tmp_path)  # >500 rows: a batch is flushed first
    with TestClient(app, raise_server_exceptions=False) as client:
        res = client.post(
            "/orders/import",
            headers=AUTH,
            files={"file": ("db-fail.xlsx", payload, XLSX_CT)},
        )
        assert res.status_code == 500
        # Starlette renders an unhandled error as plain text; the point is that
        # the client is not told the workbook was malformed.
        assert "invalid xlsx" not in res.text
        listed = client.get("/orders", headers=AUTH)
    assert listed.status_code == 200
    assert listed.json() == []
    assert rolled_back == ["rollback"]


def test_customer_import_discards_stale_synthetic_meta(demo_client, tmp_path, monkeypatch) -> None:
    """M-8: a label may only survive for the workbook that produced it."""
    seed_dir = tmp_path / "seed"
    monkeypatch.setattr(dbmod, "SEED_DIR", seed_dir)
    meta = generate(20, seed=31, out_dir=seed_dir)
    meta_path = seed_dir / dbmod.DATASET_META_NAME
    payload = (seed_dir / str(meta["xlsx"])).read_bytes()

    assert _post(demo_client, payload).status_code == 200
    assert meta_path.exists()
    assert _post(demo_client, payload).json()["rows_imported"] == 0  # all duplicates
    assert meta_path.exists()

    other = _legacy_xlsx([["Q1", 10.776, 106.7, "KH", "0900000002", 8, "08:00", "10:00", "thuong", ""]])
    assert _post(demo_client, other, "customer.xlsx").status_code == 200
    assert not meta_path.exists()


def test_reseed_discards_stale_synthetic_meta(demo_client, tmp_path, monkeypatch) -> None:
    seed_dir = tmp_path / "seed"
    monkeypatch.setattr(dbmod, "SEED_DIR", seed_dir)
    generate(20, seed=32, out_dir=seed_dir)
    meta_path = seed_dir / dbmod.DATASET_META_NAME
    assert meta_path.exists()
    assert demo_client.post("/seed", headers=AUTH).status_code == 200
    assert not meta_path.exists()


def test_external_id_gets_column_and_index_on_legacy_db(tmp_path) -> None:
    """M-12: fresh DBs index the dedupe key; migrated ones must too."""
    with dbmod.engine.connect() as conn:
        fresh = {row[1] for row in conn.exec_driver_sql("PRAGMA index_list(orders)").fetchall()}
    assert "ix_orders_external_id" in fresh

    legacy = tmp_path / "legacy.db"
    con = sqlite3.connect(legacy)
    con.execute("CREATE TABLE orders (id INTEGER PRIMARY KEY, address TEXT)")
    con.commit()
    con.close()
    dbmod.set_engine(f"sqlite:///{legacy}")
    with dbmod.engine.connect() as conn:
        cols = {row[1] for row in conn.exec_driver_sql("PRAGMA table_info(orders)").fetchall()}
        index = {row[1] for row in conn.exec_driver_sql("PRAGMA index_list(orders)").fetchall()}
    assert "external_id" in cols
    assert "ix_orders_external_id" in index


def test_row_cap_boundary_exact_and_plus_one(demo_client, monkeypatch, tmp_path) -> None:
    """N-13: exactly max_data_rows imports; one more rolls the whole file back."""
    monkeypatch.setenv("ECOMILES_MAX_DATA_ROWS", "500")
    exact = _post(demo_client, _generated(500, 19, tmp_path / "exact"), "exact.xlsx")
    assert exact.status_code == 200
    body = exact.json()
    assert body["rows_seen"] == 500
    assert body["rows_imported"] == 500
    assert body["imported"] == 500
    assert body["skipped_duplicates"] == 0
    assert body["rows_rejected"] == 0
    assert len(demo_client.get("/orders", headers=AUTH).json()) == 500

    over = _post(demo_client, _generated(501, 23, tmp_path / "over"), "over.xlsx")
    assert over.status_code == 400
    assert over.json()["detail"] == "too many rows"
    # the earlier import is untouched: the rejected file committed nothing
    assert len(demo_client.get("/orders", headers=AUTH).json()) == 500


def test_stream_counters_reset_on_reuse(tmp_path) -> None:
    """M-12: iterating the same stream twice must not double-count."""
    stream = IngestStream(xlsx_bytes(order_rows(3, 29)))
    first = list(stream)
    assert (stream.rows_seen, len(first), stream.dedupe) == (3, 3, True)
    second = list(stream)
    assert (stream.rows_seen, len(second)) == (3, 3)
    assert stream.skipped_duplicates == 0
