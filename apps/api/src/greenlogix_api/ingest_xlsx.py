"""openpyxl xlsx ingest with bilingual headers (ORD-01, C-03). No geocoding."""

from __future__ import annotations

import logging
import os
import re
from collections.abc import Callable, Iterable, Iterator
from datetime import datetime, time
from io import BytesIO

from openpyxl import load_workbook

from greenlogix_api.models import Order
from greenlogix_api.schemas import CARGO_TYPES, LOCAL_TIME_PATTERN, ImportErrorItem

log = logging.getLogger("greenlogix")

DEFAULT_MAX_UPLOAD_BYTES = 33_554_432  # 32 MB
DEFAULT_MAX_DATA_ROWS = 50_000
# Hard ceiling for ECOMILES_MAX_UPLOAD_BYTES: 256 MB. Anything above is clamped
# (and logged) so a typo cannot turn the ingest path into an unbounded buffer.
MAX_UPLOAD_BYTES_CEILING = 268_435_456
_WINDOW_RE = re.compile(LOCAL_TIME_PATTERN)


def _env_int(name: str, default: int, *, ceiling: int | None = None) -> int:
    raw = os.environ.get(name)
    if raw is None:
        return default
    try:
        value = int(raw)
    except ValueError:
        value = 0
    if value <= 0:
        log.warning("ignoring %s=%r, not a positive integer", name, raw)
        return default
    if ceiling is not None and value > ceiling:
        log.warning("clamping %s=%s to ceiling %s", name, value, ceiling)
        return ceiling
    return value


def resolve_limits() -> tuple[int, int]:
    """(max_upload_bytes, max_data_rows), read from env on every call.

    ``ECOMILES_MAX_UPLOAD_BYTES`` / ``ECOMILES_MAX_DATA_ROWS``; invalid or
    non-positive values fall back to the raised defaults, and the byte cap is
    clamped to ``MAX_UPLOAD_BYTES_CEILING`` (256 MB).
    """
    return (
        _env_int(
            "ECOMILES_MAX_UPLOAD_BYTES", DEFAULT_MAX_UPLOAD_BYTES, ceiling=MAX_UPLOAD_BYTES_CEILING
        ),
        _env_int("ECOMILES_MAX_DATA_ROWS", DEFAULT_MAX_DATA_ROWS),
    )


# Backwards-compatible constant names: import-time snapshot of the same caps.
MAX_UPLOAD_BYTES, MAX_DATA_ROWS = resolve_limits()

HEADER_ALIASES: dict[str, frozenset[str]] = {
    "address": frozenset({"address", "địa chỉ"}),
    "lat": frozenset({"lat", "vĩ độ"}),
    "lng": frozenset({"lng", "kinh độ"}),
    "receiver": frozenset({"receiver", "người nhận"}),
    "phone": frozenset({"phone", "sđt"}),
    "kg": frozenset({"kg", "khối lượng"}),
    "window_start": frozenset({"window_start", "giờ bắt đầu"}),
    "window_end": frozenset({"window_end", "giờ kết thúc"}),
    "cargo_type": frozenset({"cargo_type", "loại hàng"}),
    "notes": frozenset({"notes", "ghi chú"}),
    # C-03: caller-supplied dedupe key (nullable; absent column => no dedupe).
    "external_id": frozenset({"external_id", "mã đơn", "order_id"}),
}


class IngestLimitError(ValueError):
    pass


def _norm_header(value: object) -> str:
    return str(value or "").strip().lower()


def map_headers(header_row: tuple[object, ...] | list[object]) -> dict[str, int]:
    found: dict[str, int] = {}
    for idx, cell in enumerate(header_row):
        key = _norm_header(cell)
        for canon, aliases in HEADER_ALIASES.items():
            if key in aliases:
                found[canon] = idx
                break
    return found


def normalize_window(value: object) -> str:
    if value is None or value == "":
        return ""
    if isinstance(value, datetime):
        return f"{value.hour:02d}:{value.minute:02d}"
    if isinstance(value, time):
        return f"{value.hour:02d}:{value.minute:02d}"
    if isinstance(value, bool):
        return ""
    if isinstance(value, (int, float)):
        minutes = round(float(value) * 24 * 60) % (24 * 60)
        hours, mins = divmod(minutes, 60)
        return f"{hours:02d}:{mins:02d}"
    text = str(value).strip()
    if "1899" in text and " " in text:
        tail = text.split(" ")[-1]
        return normalize_window(tail)
    if len(text) >= 5 and text[2] == ":":
        return text[:5]
    return text


def _cell(row: tuple[object, ...], index: int | None) -> object:
    if index is None or index >= len(row):
        return None
    return row[index]


def _as_float(value: object) -> float | None:
    if value is None or value == "":
        return None
    try:
        return float(value)  # type: ignore[arg-type]
    except (TypeError, ValueError):
        return None


class IngestStream:
    """Streaming parse of one upload (C-03). Iterate to pull importable ``Order`` rows.

    ``read_only`` iteration keeps a 72k-row worksheet off the heap; counters fill
    in as iteration proceeds:

    - ``rows_seen``: non-blank data rows enumerated (blank rows ignored, as before)
    - ``errors`` / ``rows_rejected``: per-row rejections, same shape as before
    - ``dedupe``: True when the upload carried an external-id column
    - ``skipped_duplicates``: rows dropped because their external id was already
      known, from the database or from earlier in the same file

    Rows without an external id are never deduped, so an upload with no such
    column imports exactly as it did before this change. Stored values are
    validated against the same ``LocalTime`` / ``CargoType`` definitions the
    response models use, because a stored value ``OrderOut`` rejects would make
    ``GET /orders`` fail for every caller (F1).
    """

    def __init__(
        self,
        data: bytes,
        *,
        max_rows: int = MAX_DATA_ROWS,
        load_known_ids: Callable[[], Iterable[str]] | None = None,
    ) -> None:
        self._data = data
        self._max_rows = max_rows
        self._load_known_ids = load_known_ids
        self._known: set[str] = set()
        self.rows_seen = 0
        self.errors: list[ImportErrorItem] = []
        self.dedupe = False
        self.skipped_duplicates = 0

    @property
    def rows_rejected(self) -> int:
        return len(self.errors)

    def __iter__(self) -> Iterator[Order]:
        # Counters reset per pass: a reused stream must not double-count.
        self.rows_seen = 0
        self.errors = []
        self.skipped_duplicates = 0
        self.dedupe = False
        self._known = set()
        wb = load_workbook(BytesIO(self._data), data_only=True, read_only=True)
        try:
            ws = wb.active
            if ws is None:
                self.errors.append(
                    ImportErrorItem(excel_row=1, field="file", message="empty workbook")
                )
                return
            rows = ws.iter_rows(values_only=True)
            headers = map_headers(tuple(next(rows, None) or ()))
            # The dedupe lookup is only worth its query when the file can carry
            # a key at all (M-12).
            self.dedupe = "external_id" in headers
            if self.dedupe and self._load_known_ids is not None:
                self._known = {eid for eid in self._load_known_ids() if eid}
            for offset, raw in enumerate(rows, start=2):
                row = tuple(raw or ())
                if all(cell is None or cell == "" for cell in row):
                    continue
                self.rows_seen += 1
                if self.rows_seen > self._max_rows:
                    raise IngestLimitError("too many rows")
                order = self._order(row, headers, offset)
                if order is not None:
                    yield order
        finally:
            wb.close()

    def _order(
        self, row: tuple[object, ...], headers: dict[str, int], offset: int
    ) -> Order | None:
        lat = _as_float(_cell(row, headers.get("lat")))
        if lat is None:
            self._reject(offset, "lat", "missing lat")
            return None
        lng = _as_float(_cell(row, headers.get("lng")))
        if lng is None:
            self._reject(offset, "lng", "missing lng")
            return None
        window_start = self._window(row, headers, offset, "window_start")
        if window_start is None:
            return None
        window_end = self._window(row, headers, offset, "window_end")
        if window_end is None:
            return None
        cargo_type = str(_cell(row, headers.get("cargo_type")) or "").strip() or "thuong"
        if cargo_type not in CARGO_TYPES:
            # Never echo the cell text: it is untrusted input and would be
            # reflected back to the client.
            self._reject(offset, "cargo_type", "unknown cargo_type")
            return None
        external_id = self._external_id(row, headers)
        if external_id is not None:
            if external_id in self._known:
                self.skipped_duplicates += 1
                log.info("import skip excel_row=%s duplicate", offset)
                return None
            self._known.add(external_id)
        return Order(
            address=str(_cell(row, headers.get("address")) or ""),
            lat=lat,
            lng=lng,
            receiver=str(_cell(row, headers.get("receiver")) or ""),
            phone=str(_cell(row, headers.get("phone")) or ""),
            kg=_as_float(_cell(row, headers.get("kg"))) or 0.0,
            window_start=window_start,
            window_end=window_end,
            cargo_type=cargo_type,
            notes=str(_cell(row, headers.get("notes")) or ""),
            excel_row=offset,
            external_id=external_id,
            status="pending",
        )

    def _window(
        self, row: tuple[object, ...], headers: dict[str, int], offset: int, field: str
    ) -> str | None:
        """Coerce a delivery-window cell to ``HH:MM``, else reject the row.

        A missing/blank/unparseable window is rejected rather than defaulted:
        ``OrderOut`` types these fields as ``LocalTime``, and inventing a window
        would fabricate a real delivery constraint (``window_end`` also drives
        ``late_risk``). Only an unpadded hour such as ``8:00`` is normalised.
        """
        text = normalize_window(_cell(row, headers.get(field)))
        candidate = f"0{text}" if re.fullmatch(r"\d:\d{2}", text) else text
        if _WINDOW_RE.fullmatch(candidate):
            return candidate
        self._reject(offset, field, f"missing {field}" if not text else f"invalid {field}")
        return None

    def _reject(self, offset: int, field: str, message: str) -> None:
        self.errors.append(ImportErrorItem(excel_row=offset, field=field, message=message))
        log.info("import skip excel_row=%s field=%s", offset, field)

    def _external_id(self, row: tuple[object, ...], headers: dict[str, int]) -> str | None:
        raw = _cell(row, headers.get("external_id"))
        if raw is None:
            return None
        text = str(raw).strip()
        return text or None
