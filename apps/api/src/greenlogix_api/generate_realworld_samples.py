"""Generate 3 near-real-world Vietnamese logistics .xlsx sample workbooks for customer & judge benchmarking."""

from __future__ import annotations

import json
import random
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

from greenlogix_api.generate_synthetic import FIXED_MTIME, _pin_zip_timestamps

# 61 Verified Real Ho Chi Minh City Commercial Street Corridors (Prime count = 61 so modular indexing
# visits ALL 61 corridors evenly across 15 HCMC districts; strictly zero points on Tân Sơn Nhất Airport!)
HCMC_CORRIDORS: list[tuple[str, str, float, float]] = [
    # Quận 1 (8 real streets)
    ("Quận 1", "Nguyễn Thị Minh Khai, P. Đa Kao", 10.7872, 106.7012),
    ("Quận 1", "Hai Bà Trưng, P. Tân Định", 10.7905, 106.6912),
    ("Quận 1", "Lê Thánh Tôn, P. Bến Nghé", 10.7778, 106.7032),
    ("Quận 1", "Lê Lợi, P. Bến Thành", 10.7726, 106.6982),
    ("Quận 1", "Nguyễn Trãi, P. Nguyễn Cư Trinh", 10.7662, 106.6875),
    ("Quận 1", "Trần Hưng Đạo, P. Cầu Kho", 10.7615, 106.6888),
    ("Quận 1", "Cô Giang, P. Cô Giang", 10.7642, 106.6938),
    ("Quận 1", "Đinh Tiên Hoàng, P. Đa Kao", 10.7898, 106.6978),
    # Quận 3 (7 real streets)
    ("Quận 3", "Võ Văn Tần, Phường 6", 10.7768, 106.6895),
    ("Quận 3", "Nam Kỳ Khởi Nghĩa, Phường 8", 10.7885, 106.6852),
    ("Quận 3", "Cách Mạng Tháng 8, Phường 10", 10.7815, 106.6748),
    ("Quận 3", "Điện Biên Phủ, Phường 4", 10.7745, 106.6812),
    ("Quận 3", "Lê Văn Sỹ, Phường 14", 10.7892, 106.6782),
    ("Quận 3", "Nguyễn Đình Chiểu, Phường 5", 10.7718, 106.6845),
    ("Quận 3", "Kỳ Đồng, Phường 9", 10.7832, 106.6808),
    # Quận 10 & Quận 11 (8 real streets)
    ("Quận 10", "Ba Tháng Hai, Phường 12", 10.7715, 106.6698),
    ("Quận 10", "Sư Vạn Hạnh, Phường 13", 10.7762, 106.6668),
    ("Quận 10", "Tô Hiến Thành, Phường 15", 10.7802, 106.6635),
    ("Quận 10", "Thành Thái, Phường 14", 10.7738, 106.6612),
    ("Quận 10", "Ngô Gia Tự, Phường 4", 10.7642, 106.6712),
    ("Quận 10", "Lý Thường Kiệt, Phường 14", 10.7728, 106.6578),
    ("Quận 10", "Nguyễn Tri Phương, Phường 8", 10.7618, 106.6675),
    ("Quận 11", "Ông Ích Khiêm, Phường 10", 10.7658, 106.6435),
    # Quận 5 & Quận 6 (8 real streets)
    ("Quận 5", "Nguyễn Trãi, Phường 8", 10.7558, 106.6725),
    ("Quận 5", "Trần Hưng Đạo, Phường 2", 10.7562, 106.6805),
    ("Quận 5", "An Dương Vương, Phường 4", 10.7602, 106.6772),
    ("Quận 5", "Hồng Bàng, Phường 11", 10.7545, 106.6615),
    ("Quận 5", "Hải Thượng Lãn Ông, Phường 10", 10.7508, 106.6602),
    ("Quận 6", "Hậu Giang, Phường 6", 10.7495, 106.6465),
    ("Quận 6", "Kinh Dương Vương, Phường 12", 10.7512, 106.6348),
    ("Quận 6", "Bình Tiên, Phường 3", 10.7468, 106.6498),
    # Tân Bình & Tân Phú (8 real streets — strictly South/West of Cộng Hòa, never on Airport)
    ("Tân Bình", "Cộng Hòa, Phường 13", 10.8012, 106.6415),
    ("Tân Bình", "Hoàng Văn Thụ, Phường 4", 10.7975, 106.6568),
    ("Tân Bình", "Lý Thường Kiệt, Phường 8", 10.7865, 106.6538),
    ("Tân Bình", "Lạc Long Quân, Phường 10", 10.7805, 106.6458),
    ("Tân Bình", "Bàu Cát, Phường 14", 10.7912, 106.6418),
    ("Tân Phú", "Lũy Bán Bích, P. Hòa Thạnh", 10.7745, 106.6318),
    ("Tân Phú", "Tân Kỳ Tân Quý, P. Sơn Kỳ", 10.8015, 106.6238),
    ("Tân Phú", "Hòa Bình, P. Hiệp Tân", 10.7685, 106.6335),
    # Phú Nhuận & Gò Vấp (7 real streets)
    ("Phú Nhuận", "Phan Xích Long, Phường 2", 10.7982, 106.6895),
    ("Phú Nhuận", "Phan Đăng Lưu, Phường 3", 10.8005, 106.6828),
    ("Phú Nhuận", "Nguyễn Văn Trỗi, Phường 11", 10.7942, 106.6738),
    ("Phú Nhuận", "Huỳnh Văn Bánh, Phường 17", 10.7918, 106.6788),
    ("Gò Vấp", "Phan Văn Trị, Phường 7", 10.8235, 106.6885),
    ("Gò Vấp", "Quang Trung, Phường 10", 10.8285, 106.6712),
    ("Gò Vấp", "Nguyễn Oanh, Phường 17", 10.8352, 106.6768),
    # Bình Thạnh (6 real streets)
    ("Bình Thạnh", "Xô Viết Nghệ Tĩnh, Phường 21", 10.7975, 106.7105),
    ("Bình Thạnh", "Điện Biên Phủ, Phường 25", 10.8012, 106.7158),
    ("Bình Thạnh", "Bạch Đằng, Phường 15", 10.8028, 106.7018),
    ("Bình Thạnh", "Nguyễn Gia Trí (D2), Phường 25", 10.8045, 106.7162),
    ("Bình Thạnh", "Lê Quang Định, Phường 14", 10.8112, 106.6952),
    ("Bình Thạnh", "Nguyễn Hữu Cảnh, Phường 22", 10.7928, 106.7188),
    # Quận 4, Quận 7 & Quận 8 (5 real streets)
    ("Quận 4", "Hoàng Diệu, Phường 6", 10.7608, 106.7015),
    ("Quận 7", "Nguyễn Thị Thập, P. Tân Phong", 10.7405, 106.7085),
    ("Quận 7", "Nguyễn Hữu Thọ, P. Tân Hưng", 10.7458, 106.7016),
    ("Quận 7", "Huỳnh Tấn Phát, P. Bình Thuận", 10.7418, 106.7262),
    ("Quận 8", "Dương Bá Trạc, Phường 1", 10.7465, 106.6878),
    # TP. Thủ Đức (4 real streets)
    ("TP. Thủ Đức", "Song Hành Xa Lộ Hà Nội, P. Thảo Điền", 10.8035, 106.7382),
    ("TP. Thủ Đức", "Trần Não, P. An Khánh", 10.7885, 106.7315),
    ("TP. Thủ Đức", "Phạm Văn Đồng, P. Hiệp Bình Chánh", 10.8275, 106.7218),
    ("TP. Thủ Đức", "Kha Vạn Cân, P. Linh Đông", 10.8385, 106.7425),
]
assert len(HCMC_CORRIDORS) == 61, f"Expected 61 prime corridors, got {len(HCMC_CORRIDORS)}"

STORE_CHAINS = [
    "Bách Hóa Xanh",
    "WinMart+",
    "Co.op Food",
    "GS25",
    "Circle K",
    "Siêu thị Ân Nam",
    "Nhà hàng Golden Gate",
    "Highlands Coffee",
    "Phúc Long Hub",
    "Đại lý Phân Phối Minh Phát",
]

HEADERS = [
    "external_id",
    "address",
    "lat",
    "lng",
    "receiver",
    "phone",
    "kg",
    "window_start",
    "window_end",
    "cargo_type",
    "notes",
]

FLEET_ROWS = [
    ["51C-000.01", "van_500", 500, "petrol", 8.0, "ready"],
    ["51C-000.02", "truck_800", 800, "diesel", 9.0, "ready"],
    ["51C-000.03", "truck_1000", 1000, "petrol", 10.0, "ready"],
    ["51C-000.04", "truck_1500", 1500, "diesel", 11.0, "ready"],
    ["51C-000.05", "truck_2000", 2000, "petrol", 12.0, "ready"],
    ["51C-000.06", "van_500", 500, "diesel", 8.5, "ready"],
    ["51C-000.07", "truck_800", 800, "petrol", 9.5, "ready"],
    ["51C-000.08", "truck_1000", 1000, "diesel", 8.0, "ready"],
    ["51C-000.09", "truck_1500", 1500, "petrol", 10.0, "ready"],
    ["51C-000.10", "truck_2000", 2000, "diesel", 12.0, "ready"],
]


def _styled_workbook(title: str, rows: list[list[object]]) -> bytes:
    from io import BytesIO

    wb = Workbook()
    ws = wb.active
    assert ws is not None
    ws.title = "orders"
    ws.append(HEADERS)

    header_fill = PatternFill(start_color="0F172A", end_color="0F172A", fill_type="solid")
    header_font = Font(bold=True, color="FFD400")
    for col_idx in range(1, len(HEADERS) + 1):
        cell = ws.cell(row=1, column=col_idx)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center")

    for row in rows:
        ws.append(row)

    ws_fleet = wb.create_sheet(title="Doi_Xe_Hien_Co")
    ws_fleet.append(["plate", "type", "capacity_kg", "fuel", "l_per_100km", "status"])
    for f_row in FLEET_ROWS:
        ws_fleet.append(f_row)

    wb.properties.title = title
    wb.properties.created = FIXED_MTIME
    wb.properties.modified = FIXED_MTIME
    buf = BytesIO()
    wb.save(buf)
    return _pin_zip_timestamps(buf.getvalue())


def build_fmcg_80(seed: int = 202601) -> list[list[object]]:
    rng = random.Random(seed)
    rows: list[list[object]] = []
    windows = [
        ("09:00", "11:30"),
        ("09:30", "12:00"),
        ("10:00", "13:00"),
        ("13:00", "15:30"),
        ("13:30", "16:00"),
    ]
    for i in range(1, 81):
        # Step by 17 (coprime to 61) so ERP export zig-zags across all 61 real HCMC streets
        dist, street, lat0, lng0 = HCMC_CORRIDORS[(i * 17) % len(HCMC_CORRIDORS)]
        chain = STORE_CHAINS[(i * 3) % len(STORE_CHAINS)]
        lat = round(lat0 + rng.uniform(-0.0018, 0.0018), 6)
        lng = round(lng0 + rng.uniform(-0.0018, 0.0018), 6)
        w_start, w_end = windows[i % len(windows)]
        kg = round(rng.uniform(18.0, 145.0), 1)
        house_no = 10 + (i * 13) % 380
        is_backhaul = i % 6 == 0
        note = (
            f"BACKHAUL_PICKUP: Thu hồi kiện hàng/bao bì chiều về kho từ {dist}"
            if is_backhaul
            else f"Kênh GT/MT {dist}; giao kiện tiêu chuẩn"
        )
        rows.append(
            [
                f"FMCG-2026-{i:04d}",
                f"{house_no} {street}, {dist}, TP.HCM",
                lat,
                lng,
                f"{chain} - CN {dist} #{i:02d}",
                f"090{rng.randrange(10**7):07d}",
                kg,
                w_start,
                w_end,
                "thuong" if i % 4 != 0 else "de_vo",
                note,
            ]
        )
    return rows


def build_coldchain_120_qd23(seed: int = 202602) -> list[list[object]]:
    rng = random.Random(seed)
    rows: list[list[object]] = []
    for i in range(1, 121):
        # Step by 19 (coprime to 61) so all 61 real HCMC corridors are covered ( ~2 stops per corridor )
        dist, street, lat0, lng0 = HCMC_CORRIDORS[(i * 19) % len(HCMC_CORRIDORS)]
        chain = STORE_CHAINS[(i * 5 + 1) % len(STORE_CHAINS)]
        lat = round(lat0 + rng.uniform(-0.0016, 0.0016), 6)
        lng = round(lng0 + rng.uniform(-0.0016, 106.75 - lng0 if lng0 > 106.74 else 0.0016), 6)
        # 28 orders intentionally scheduled inside QĐ 23/2018 truck ban windows (06:00-09:00 or 16:00-20:00)
        if i <= 16:
            w_start, w_end = "06:30", "08:30"
            note = "CẢNH BÁO QĐ 23/2018: Khung giờ cao điểm sáng (06:00-09:00); Hàng lạnh -18C"
        elif i <= 28:
            w_start, w_end = "16:30", "18:30"
            note = "CẢNH BÁO QĐ 23/2018: Khung giờ cao điểm chiều (16:00-20:00); Hàng tươi sống"
        elif i % 6 == 0:
            w_start, w_end = "14:00", "15:50"
            note = f"BACKHAUL_PICKUP: Nhận hàng thu hồi chiều về kho từ {dist}"
        elif i % 3 == 0:
            w_start, w_end = "09:15", "11:30"
            note = "Hàng lạnh 2-4C; ưu tiên dỡ hàng nhanh"
        else:
            w_start, w_end = "12:30", "15:30"
            note = "Chuỗi lạnh F&B; kiểm nhiệt độ khi nhận"
        kg = round(rng.uniform(22.0, 160.0), 1)
        house_no = 5 + (i * 17) % 420
        rows.append(
            [
                f"COLD-2026-{i:04d}",
                f"{house_no} {street}, {dist}, TP.HCM",
                lat,
                lng,
                f"{chain} ColdHub #{i:03d}",
                f"091{rng.randrange(10**7):07d}",
                kg,
                w_start,
                w_end,
                "lanh" if i % 5 != 0 else "de_vo",
                note,
            ]
        )
    return rows


def build_erp_500_multiday(seed: int = 202603) -> list[list[object]]:
    rng = random.Random(seed)
    rows: list[list[object]] = []
    days = [f"2026-10-{18 + d:02d}" for d in range(7)]
    windows = [
        ("07:30", "09:30"),
        ("09:00", "11:00"),
        ("10:00", "12:30"),
        ("13:00", "15:30"),
        ("14:00", "16:30"),
        ("16:30", "19:00"),
    ]
    for i in range(1, 501):
        day = days[(i - 1) % len(days)]
        dist, street, lat0, lng0 = HCMC_CORRIDORS[(i * 23) % len(HCMC_CORRIDORS)]
        chain = STORE_CHAINS[(i * 7) % len(STORE_CHAINS)]
        lat = round(lat0 + rng.uniform(-0.0022, 0.0022), 6)
        lng = round(lng0 + rng.uniform(-0.0022, 0.0022), 6)
        w_start, w_end = windows[(i - 1) % len(windows)]
        kg = round(rng.uniform(12.0, 175.0), 1)
        cargo = "lanh" if i % 4 == 0 else ("de_vo" if i % 7 == 0 else "thuong")
        house_no = 1 + (i * 19) % 500
        note = (
            f"date={day}; BACKHAUL_PICKUP=1; district={dist}"
            if i % 6 == 0
            else f"date={day}; ERP_BATCH=DN_A_W43; district={dist}"
        )
        rows.append(
            [
                f"DNA-{day}-{i:04d}",
                f"{house_no} {street}, {dist}, TP.HCM",
                lat,
                lng,
                f"{chain} ({dist}) #{i:03d}",
                f"098{rng.randrange(10**7):07d}",
                kg,
                w_start,
                w_end,
                cargo,
                note,
            ]
        )
    return rows


def generate_all(out_dir: Path) -> dict[str, object]:
    out_dir.mkdir(parents=True, exist_ok=True)
    datasets = {
        "EcoMiles_Mau_FMCG_80_Don_TPHCM.xlsx": (
            "FMCG 80 đơn nội đô TP.HCM (Kịch bản điều phối thủ công vs EcoMiles)",
            build_fmcg_80(),
        ),
        "EcoMiles_Mau_Chuoi_Lanh_120_Don_Cam_Tai_QD23.xlsx": (
            "Cold-Chain 120 đơn F&B (28 đơn vi phạm giờ cấm tải QĐ 23/2018)",
            build_coldchain_120_qd23(),
        ),
        "EcoMiles_Mau_ERP_500_Don_7_Ngay.xlsx": (
            "ERP Doanh Nghiệp A — 500 đơn hàng đa ngày (7 ngày vận hành)",
            build_erp_500_multiday(),
        ),
    }
    manifest: dict[str, object] = {}
    for fname, (title, rows) in datasets.items():
        raw = _styled_workbook(title, rows)
        (out_dir / fname).write_bytes(raw)
        manifest[fname] = {"title": title, "rows": len(rows), "bytes": len(raw)}
    (out_dir / "samples_manifest.json").write_text(
        json.dumps(manifest, indent=2, ensure_ascii=False) + "\n", encoding="utf-8"
    )
    return manifest


if __name__ == "__main__":
    root = Path(__file__).resolve().parents[4]
    m1 = generate_all(root / "apps" / "landing" / "public" / "samples")
    print(json.dumps(m1, indent=2, ensure_ascii=False))
